"""CSM draft renderer: the working draft a CSM edits and takes into the customer conversation, built from
output/<company>/plan.json.

No model call and no new content: every line comes from plan.json (or from the capability display names in
config/draft_tasks.yaml), apart from fixed headings and "CSM action" instructions. Writes csm_draft.md and
csm_draft.docx next to the plan, both built from one list of blocks so they say the same thing. The CSM draft carries
no sources or citations (those are in plan.md) but keeps the provenance and fit labels. Discovery questions sit
under the objective, criterion, or challenge they belong to, each shown once. Only challenges on the keep lists in
config/csm_draft_rules.yaml appear; the rest, criterion validation questions, and questions about initiatives the tool did
not map stay in plan.md. The Render leakage checks run on the CSM draft before anything is kept.
"""

import json
import re
from pathlib import Path

import yaml
from docx import Document
from docx.shared import Inches, Pt

from .. import atomic, guards
from ..models import PLAN_CONFIRMATION_BLANK
from ..paths import PROJECT_ROOT, CompanyPaths
from .draft import FIT_LABELS, TASKS_FILE
from .render import (
    DISCLAIMER, OMITTED_ITEM_TEXT, OMITTED_MARKER, PLAN_JSON, RenderError, STATUS, evidence_problems,
    load_chunk_texts, output_dir,
)

DRAFT_MD, DRAFT_DOCX = "csm_draft.md", "csm_draft.docx"
RULES_FILE = PROJECT_ROOT / "config" / "csm_draft_rules.yaml"
TITLE = "Success Plan Strategy and Initiatives"
BANNER = "Draft hypothesis from public information. Not reviewed by a human. For CSM review."
POINTER = "Sources, citations and evidence labels are in the detailed plan (plan.md)."
SECTIONS = [
    "1. Company Highlights",
    "2. Key Business Initiatives",
    "3. Key Benefits",
    "4. Renewal",
    "5. Success Criteria",
    "6. Key Challenges / Potential Blockers",
    "7. Objectives, Success Criteria, Tasks",
]
ACTIONS = [
    "Add the account history and confirm the overview matches what the customer says.",
    "Select the initiatives relevant to this account, from the bullets and from the no-fit line. Objectives below are "
    "drafted for every initiative with a fit; remove the ones not selected.",
    "Confirm which plan includes these capabilities and which benefits matter to the customer.",
    "Complete the blank.",
    "Each criterion names what to measure and which direction is good. Get the customer's current baseline in "
    "discovery, then set the target with the customer.",
    "Validate which of these are relevant to this account and add any the customer raises.",
    "Keep what applies, answer the questions under each objective in discovery, and agree baselines before setting "
    "targets.",
]
ACTION_LABEL = "CSM action:"
TARGET_BLANK_ONCE = "[CSM to complete: set target after baseline]"
CUSTOMER_RAISED = "Raised by the customer: [CSM to complete]"
ENGAGEMENT_GROUP = "Could affect this engagement"
DESCRIPTION_OMITTED = "[Description omitted in the CSM draft; see the detailed plan]"
NO_FIT_PREFIX = "No product fit: "
BASELINES_ONCE = "Agree baselines for all success criteria before setting targets (CSM, Day 31-60, proposed)"
VALIDATE, BASELINE = "Validate: ", "Baseline: "

# Block styles map one-to-one to Word styles (ACTION is a Normal paragraph with a bold label).
TITLE_S, SUBTITLE_S, NORMAL, H1, H2, ACTION, BULLET, BULLET2, BULLET3 = (
    "Title", "Subtitle", "Normal", "Heading 1", "Heading 2", "Action", "List Bullet", "List Bullet 2", "List Bullet 3")
LIST_STYLES = (BULLET, BULLET2, BULLET3)
REF = re.compile(r"\[(\d+)\]")


def load_rules(path: Path = RULES_FILE) -> dict[str, set[str]]:
    data = yaml.safe_load(path.read_text()) or {}
    return {"categories": set(data.get("challenge_categories_with_question") or []),
            "change_types": set(data.get("change_types_with_question") or [])}


def extract_fields(paths: CompanyPaths) -> dict[str, dict]:
    """category and change_type by item id, from 02_extract.json (plan.json does not carry them)."""
    path = paths.stages / "02_extract.json"
    if not path.exists():
        return {}
    data = json.loads(path.read_text())
    return {i["id"]: {"category": i.get("category"), "change_type": i.get("change_type")}
            for k in ("challenges", "changes") for i in data.get(k) or []}


def keeps_question(fields: dict, rules: dict[str, set[str]]) -> bool:
    """True if the item's category or change type is on a keep list; anything else (including unknown) is context."""
    return (fields.get("category") in rules["categories"]) or (fields.get("change_type") in rules["change_types"])


def _shown_sentence(item: dict) -> str | None:
    """First sentence of an item that is shown in the plan and states no currency amount (company financial
    figures stay in the detailed plan)."""
    for s in item.get("sentences") or []:
        if s["text"] != OMITTED_MARKER and not guards.find_currency_amounts(s["text"]):
            return s["text"]
    text = item.get("text", "")
    if text and text != OMITTED_ITEM_TEXT and not item.get("sentences") and not guards.find_currency_amounts(text):
        return text
    return None


def _measure(c: dict) -> str:
    """Direction and measure, from the proposed target ("Increase: <measure>. Target: [CSM ...]") that Draft builds;
    the target blank itself is shown once above the list. A criterion without one shows its text (a CSM blank)."""
    target = c.get("proposed_target") or ""
    m = re.match(r"(Increase|Decrease|Reach): (.+?)\. Target: ", target)
    return f"{m.group(1)}: {m.group(2)}" if m else c["text"]


def baseline_task_text() -> str:
    """Text of the code-added "agree baselines" task, from the task catalog."""
    tasks = (yaml.safe_load(TASKS_FILE.read_text()) or {}).get("tasks") or []
    return next((t["text"] for t in tasks if t.get("id") == "T-BASELINES"), "")


def others_line(n: int) -> str:
    if n == 1:
        return "1 other public risk or change is listed in the detailed plan (plan.md)."
    return f"{n} other public risks and changes are listed in the detailed plan (plan.md)."


def _tags(*parts) -> str:
    tags = [p for p in parts if p]
    return f" ({'; '.join(tags)})" if tags else ""


def build_blocks(plan: dict, display_names: dict[str, str], fields: dict[str, dict] | None = None,
                 rules: dict[str, set[str]] | None = None) -> list[tuple[str, str]]:
    m = plan["metadata"]
    initiatives = {i["id"]: i for i in plan["key_business_initiatives"]}
    benefits = {b["initiative_id"]: b for b in plan["key_benefits"]}
    statuses = {i["id"]: STATUS[i.get("csm_selected")] for i in plan["key_business_initiatives"]}
    shown_questions: set[str] = set()

    def question(prefix: str, q: str | None, style: str) -> list[tuple[str, str]]:
        """A question line, once per CSM draft."""
        if not q or q in shown_questions:
            return []
        shown_questions.add(q)
        return [(style, prefix + q)]

    def section(n: int) -> list[tuple[str, str]]:
        return [(H1, SECTIONS[n]), (ACTION, ACTIONS[n])]

    B: list[tuple[str, str]] = [(TITLE_S, TITLE), (SUBTITLE_S, m["company"]), (NORMAL, BANNER), (NORMAL, DISCLAIMER),
                                (NORMAL, POINTER)]

    # 1. Company Highlights
    B += section(0)
    B += [(NORMAL, plan["company_highlights"]["overview"]["text"]),
          (NORMAL, f"Account history: {plan['company_highlights']['account_history']['placeholder']}")]

    # 2. Key Business Initiatives
    B += section(1)
    same_status = len(set(statuses.values())) == 1
    if same_status:
        B.append((NORMAL, f"All initiatives: {next(iter(statuses.values()))}."))
    fit_ids = {b["initiative_id"] for b in plan["key_benefits"] if b["fit"] != "no_fit"}
    for i in (x for x in plan["key_business_initiatives"] if x["id"] in fit_ids):
        shown = _shown_sentence(i)
        if not shown:
            B.append((BULLET, f"{i['title']}: {DESCRIPTION_OMITTED}"))
            continue
        B.append((BULLET, f"{i['title']}: {shown}" + _tags(None if same_status else statuses[i["id"]], i.get("provenance"),
                                                         *(i.get("labels") or []))))

    no_fit_titles = [x["title"] for x in plan["key_business_initiatives"] if x["id"] not in fit_ids]
    if no_fit_titles:
        B.append((NORMAL, NO_FIT_PREFIX + "; ".join(no_fit_titles)))

    # 3. Key Benefits
    B += section(2)
    fitted = [b for b in plan["key_benefits"] if b["fit"] != "no_fit"]
    for b in fitted:
        names = [display_names.get(c, c) for c in b.get("capability_ids") or []]
        title = initiatives.get(b["initiative_id"], {}).get("title", b["initiative_id"])
        B.append((BULLET, f"{title} ({b['initiative_id']}): {', '.join(names)}"
                          + _tags(FIT_LABELS.get(b["fit"]), b.get("provenance"), *(b.get("labels") or []))))
    n_no_fit = len(plan["key_benefits"]) - len(fitted)
    B.append((NORMAL, f"No fit for {n_no_fit} initiative{'' if n_no_fit == 1 else 's'} (listed in the detailed plan)."))
    if fitted:
        B.append((NORMAL, PLAN_CONFIRMATION_BLANK))

    # 4. Renewal
    B += section(3) + [(NORMAL, plan["renewal"]["placeholder"])]

    # 5. Success Criteria: measures only (direction and measure); questions are under the objectives
    B += section(4) + [(NORMAL, TARGET_BLANK_ONCE)]
    for c in plan["success_criteria"]:
        kind = ", ".join(x.replace("_", " ") for x in (c.get("criterion_type"), c.get("indicator_type")) if x)
        measure = _measure(c)
        B.append((BULLET, f"{c['id']} {measure}" + _tags(kind, c.get("provenance"), *(c.get("labels") or []))))

    # 6. Key Challenges / Potential Blockers
    B += section(5)
    fields, rules = fields or {}, rules or {"categories": set(), "change_types": set()}
    signals = plan["key_challenges"]["public_signals"]  # sentences are in the detailed plan
    engaged = [c for c in signals if keeps_question(fields.get(c["id"], {}), rules)]
    context = [c for c in signals if not keeps_question(fields.get(c["id"], {}), rules)]
    B.append((H2, ENGAGEMENT_GROUP))
    for c in engaged:
        B.append((BULLET, c["title"] + _tags(c.get("provenance"), *(c.get("labels") or []))))
        B += question(VALIDATE, c.get("validation_question"), BULLET2)
    if not engaged:
        B.append((NORMAL, "None."))
    B.append((NORMAL, others_line(len(context))))  # the other items, with their questions, are in plan.md only
    B.append((NORMAL, CUSTOMER_RAISED))

    # 7. Objectives, Success Criteria, Tasks
    B += section(6)
    baselines_text = baseline_task_text()
    if any(t["text"] == baselines_text for ob in plan["objectives"] for t in ob["tasks"]):
        B.append((NORMAL, BASELINES_ONCE))  # once here instead of under every objective (plan.md keeps each copy)
    for ob in plan["objectives"]:
        B.append((BULLET, f"{ob['id']} {ob['title']}" + _tags(*(ob.get("labels") or []))))
        for iid in ob["linked_initiative_ids"]:
            B += question(VALIDATE, initiatives.get(iid, {}).get("validation_question"), BULLET2)
            B += question(VALIDATE, benefits.get(iid, {}).get("validation_question"), BULLET2)
        for c in ob["success_criteria"]:
            B.append((BULLET2, f"Criterion {c['id']}: {c['text']}"))
            B += question(BASELINE, c.get("baseline_question"), BULLET3)  # its validation question is in plan.md
        for t in (t for t in ob["tasks"] if t["text"] != baselines_text):
            window = t["proposed_timeline"]["window"]
            B.append((BULLET2, f"Task: {t['text']} ({t.get('responsible_role', '')}, {window}, proposed)"))
    B.append((NORMAL, plan["objectives_csm_blank"]["placeholder"]))
    return B


def to_markdown(blocks: list[tuple[str, str]]) -> str:
    lines, prev = [], None
    for style, text in blocks:
        if style in LIST_STYLES and prev not in LIST_STYLES:
            lines.append("")
        if style == TITLE_S:
            lines.append(f"# {text}")
        elif style == SUBTITLE_S:
            lines += ["", f"**{text}**"]
        elif style == H1:
            lines += ["", f"## {text}"]
        elif style == H2:
            lines += ["", f"### {text}"]
        elif style == ACTION:
            lines += ["", f"**{ACTION_LABEL}** {text}"]
        elif style == BULLET:
            lines.append(f"- {text}")
        elif style == BULLET2:
            lines.append(f"  - {text}")
        elif style == BULLET3:
            lines.append(f"    - {text}")
        else:
            lines += ["", f"_{text}_" if text in (BANNER, DISCLAIMER, POINTER) else text]
        prev = style
    return "\n".join(lines).strip() + "\n"


def to_docx(blocks: list[tuple[str, str]], path: Path) -> None:
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(1)
        section.left_margin = section.right_margin = Inches(1)
    for style, text in blocks:
        if style == ACTION:
            p = doc.add_paragraph(style=NORMAL)
            p.add_run(ACTION_LABEL).bold = True
            p.add_run(f" {text}")
            continue
        p = doc.add_paragraph(text, style=style)
        if text in (BANNER, DISCLAIMER, POINTER):
            for run in p.runs:
                run.italic = True
                run.bold = text == BANNER
    atomic.save_with(path, doc.save)


def docx_text(path: Path) -> str:
    return "\n".join(p.text for p in Document(path).paragraphs)


def write_csm_draft(paths: CompanyPaths, headlines: list[str] | None = None) -> tuple[Path, Path]:
    out = output_dir(paths)
    plan_path = out / PLAN_JSON
    if not plan_path.exists():
        raise RenderError("Run `spg render` first: plan.json is missing.")
    plan = json.loads(plan_path.read_text())
    display_names = (yaml.safe_load(TASKS_FILE.read_text()) or {}).get("display_names") or {}
    blocks = build_blocks(plan, display_names, extract_fields(paths), load_rules())
    md = to_markdown(blocks)
    chunks = load_chunk_texts(paths)
    problems = draft_problems(md, chunks, headlines or [])
    if problems:
        raise RenderError("CSM draft check failed (nothing written):\n" + "\n".join(problems))
    atomic.write_text(out / DRAFT_MD, md)
    to_docx(blocks, out / DRAFT_DOCX)
    problems = draft_problems(docx_text(out / DRAFT_DOCX), chunks, headlines or [])
    if problems:
        (out / DRAFT_MD).unlink()
        (out / DRAFT_DOCX).unlink()
        raise RenderError("CSM draft check failed on the .docx (nothing kept):\n" + "\n".join(problems))
    return out / DRAFT_MD, out / DRAFT_DOCX


def draft_problems(text: str, chunk_texts: list[str] | None, headlines: list[str]) -> list[str]:
    """The Render leakage checks (no evidence field, no 15+ word source run, no headline), plus no numbered source
    reference, price, plan tier, personal title, or customer framing in the CSM draft."""
    problems = evidence_problems({}, text, chunk_texts, headlines)
    ref = REF.search(text)
    if ref:
        problems.append(f"CSM draft contains a numbered source reference ({ref.group(0)!r})")
    if guards.find_currency_amounts(text):
        problems.append("CSM draft states a currency amount")
    tier = guards.plan_tier_phrase(text.replace(PLAN_CONFIRMATION_BLANK, ""))
    if tier:
        problems.append(f"CSM draft states a plan tier ({tier!r})")
    if guards.has_personal_title(text):
        problems.append("CSM draft contains a personal title (Mr., Ms., Dr. ...)")
    framing = guards.implies_customer(text)
    if framing:
        problems.append(f"CSM draft implies the company is a customer ({framing!r})")
    return problems
