"""No tracked file, CSM input, or rendered output may contain a name from the gitignored config/private_names.txt.

The failure message lists file and line only, never the name itself.
"""

import subprocess
from pathlib import Path

import pytest

from success_plan.config import PRIVATE_NAMES_FILE, load_private_names
from success_plan.stages.retrieve.news import name_pattern

ROOT = Path(__file__).resolve().parents[1]


def tracked_files() -> list[Path]:
    out = subprocess.run(["git", "ls-files", "-z"], cwd=ROOT, capture_output=True, check=True).stdout
    return [ROOT / p for p in out.decode().split("\0") if p]


def name_hits(text: str, patterns) -> list[int]:
    """Line numbers (1-based) where any pattern matches."""
    return [n for n, line in enumerate(text.splitlines(), 1) if any(p.search(line) for p in patterns)]


def untracked_scope(root: Path) -> list[Path]:
    """Everything under data/*/inputs/ and output/, whether tracked or not (rendered plans are not committed yet)."""
    return [p for d in [*root.glob("data/*/inputs"), root / "output"] if d.is_dir() for p in d.rglob("*") if p.is_file()]


def file_text(path: Path) -> str:
    """Text of a file; for a Word document, the text of its paragraphs (a .docx is a zip, not text)."""
    if path.suffix == ".docx":
        from docx import Document
        return "\n".join(p.text for p in Document(path).paragraphs)
    return path.read_text(encoding="utf-8")


def scanned_files() -> list[Path]:
    return list(dict.fromkeys(tracked_files() + untracked_scope(ROOT)))


def test_no_private_names_in_tracked_files_inputs_or_output():
    names = load_private_names()
    if not names:
        pytest.skip("config/private_names.txt not present (it is gitignored)")
    patterns = [name_pattern(n) for n in names]
    hits = []
    for path in scanned_files():
        if path.resolve() == PRIVATE_NAMES_FILE.resolve() or not path.is_file():
            continue
        try:
            text = file_text(path)
        except UnicodeDecodeError:
            continue  # binary file
        hits += [f"{path.relative_to(ROOT)}:{n}" for n in name_hits(text, patterns)]
    assert not hits, "A private name appears in tracked files, inputs, or output (name not shown):\n" + "\n".join(hits)


def test_private_names_file_is_not_tracked():
    assert PRIVATE_NAMES_FILE not in [p.resolve() for p in tracked_files()]


def test_hits_report_lines_not_names():
    patterns = [name_pattern("Pat Example")]
    assert name_hits("intro\nquote from Pat Example here\nslug pat-example-says\nend", patterns) == [2, 3]


def test_scan_covers_inputs_and_output(tmp_path):
    (tmp_path / "output" / "testco").mkdir(parents=True)
    (tmp_path / "output" / "testco" / "plan.md").write_text("x")
    (tmp_path / "data" / "testco" / "inputs").mkdir(parents=True)
    (tmp_path / "data" / "testco" / "inputs" / "csm_selection.yaml").write_text("x")
    (tmp_path / "data" / "testco" / "raw").mkdir(parents=True)
    (tmp_path / "data" / "testco" / "raw" / "chunks.jsonl").write_text("x")
    found = {p.relative_to(tmp_path).as_posix() for p in untracked_scope(tmp_path)}
    assert found == {"output/testco/plan.md", "data/testco/inputs/csm_selection.yaml"}


def test_docx_text_is_scanned(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("A line naming Pat Example.")
    doc.save(tmp_path / "draft.docx")
    text = file_text(tmp_path / "draft.docx")
    assert name_hits(text, [name_pattern("Pat Example")]) == [1]
    (tmp_path / "output" / "testco").mkdir(parents=True)
    doc.save(tmp_path / "output" / "testco" / "csm_draft.docx")
    assert [p.name for p in untracked_scope(tmp_path)] == ["csm_draft.docx"]
