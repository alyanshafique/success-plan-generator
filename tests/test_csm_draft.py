"""CSM draft renderer tests. No network or API calls: the plan comes from the Render test fixture (fake models)."""

import json
import re

import pytest
from docx import Document

from success_plan import guards
from success_plan.models import OBJECTIVES_BLANK, PLAN_CONFIRMATION_BLANK
from success_plan.stages import csm_draft as cdr
from success_plan.stages import render as rd
from test_render import CHUNK_TEXT, HEADLINE, paths, rendered  # noqa: F401  (paths is a fixture)


def distinct_questions(paths):
    """The fixture gives every item the same questions; make them distinct so placement can be checked."""
    extract = json.loads((paths.stages / "02_extract.json").read_text())
    for k in ("initiatives", "challenges", "changes"):
        for i in extract[k]:
            i["validation_question"] = f"What is the status of {i['id']}?"
    (paths.stages / "02_extract.json").write_text(json.dumps(extract))
    mapping = json.loads((paths.stages / "03_map.json").read_text())
    for b in mapping["key_benefits"]:
        b["validation_question"] = f"Does {b['id']} fit?"
    (paths.stages / "03_map.json").write_text(json.dumps(mapping))


@pytest.fixture
def draft(paths):
    distinct_questions(paths)
    plan, plan_md, _ = rendered(paths)
    out = rd.output_dir(paths)
    return {"plan": plan, "plan_md": plan_md, "md": (out / cdr.DRAFT_MD).read_text(), "docx": out / cdr.DRAFT_DOCX,
            "paths": paths}


def section_text(md: str, n: int) -> str:
    start = md.index(f"## {cdr.SECTIONS[n]}")
    end = md.index(f"## {cdr.SECTIONS[n + 1]}") if n + 1 < len(cdr.SECTIONS) else len(md)
    return md[start:end]


def question_lines(md: str) -> list[str]:
    return [l.strip()[2:] for l in md.splitlines() if re.match(r"\s*- (Validate|Baseline): ", l)]


def test_render_writes_the_csm_draft_next_to_the_plan(draft):
    out = rd.output_dir(draft["paths"])
    assert (out / cdr.DRAFT_MD).exists() and (out / cdr.DRAFT_DOCX).exists()
    assert (out / rd.PLAN_MD).read_text() == draft["plan_md"]  # the detailed plan is unchanged


def test_title_banner_disclaimer_pointer_and_section_order(draft):
    md = draft["md"]
    assert md.startswith(f"# {cdr.TITLE}\n\n**Test Co**\n")
    assert md.index(cdr.BANNER) < md.index(rd.DISCLAIMER) < md.index(cdr.POINTER) < md.index(f"## {cdr.SECTIONS[0]}")
    assert cdr.POINTER == "Sources, citations and evidence labels are in the detailed plan (plan.md)."
    assert re.findall(r"^## (.+)$", md, re.M) == cdr.SECTIONS
    assert "Sources" not in re.findall(r"^## (.+)$", md, re.M) and "Discovery questions" not in md


def test_csm_action_line_directly_under_each_section_heading(draft):
    lines = [l for l in draft["md"].splitlines() if l.strip()]
    for n, heading in enumerate(cdr.SECTIONS):
        i = lines.index(f"## {heading}")
        assert lines[i + 1] == f"**CSM action:** {cdr.ACTIONS[n]}"
    assert draft["md"].count("**CSM action:**") == len(cdr.SECTIONS)


def test_docx_opens_with_word_styles_and_bold_action_labels(draft):
    doc = Document(draft["docx"])
    paras = doc.paragraphs
    assert paras[0].style.name == "Title" and paras[0].text == cdr.TITLE
    assert paras[1].style.name == "Subtitle" and paras[1].text == "Test Co"
    headings = [i for i, p in enumerate(paras) if p.style.name == "Heading 1"]
    assert [paras[i].text for i in headings] == cdr.SECTIONS
    for n, i in enumerate(headings):
        action = paras[i + 1]
        assert action.text == f"{cdr.ACTION_LABEL} {cdr.ACTIONS[n]}" and action.runs[0].bold and action.runs[0].text == cdr.ACTION_LABEL
    styles = {p.style.name for p in paras}
    assert styles <= {"Title", "Subtitle", "Normal", "Heading 1", "Heading 2", "List Bullet", "List Bullet 2", "List Bullet 3"}
    assert [p.text for p in paras if p.style.name == "Heading 2"] == [cdr.ENGAGEMENT_GROUP]
    assert not doc.tables
    assert doc.styles["Normal"].font.size.pt == 10
    texts = [p.text for p in paras]
    assert cdr.BANNER in texts and rd.DISCLAIMER in texts and cdr.POINTER in texts


def test_no_source_references_or_sources_section(draft):
    md, docx_text = draft["md"], cdr.docx_text(draft["docx"])
    for text in (md, docx_text):
        assert not cdr.REF.search(text) and "http" not in text
    assert any("numbered source reference" in p for p in cdr.draft_problems("A line [3].", [], []))


def test_provenance_and_fit_labels_kept(draft):
    md = draft["md"]
    assert "news headline only" in md and "(inferred" in md and "(sourced" in md
    assert "partial fit: supports the work, does not deliver the outcome" in section_text(md, 2)
    assert "company measure; not attributable to the product" in section_text(md, 4)


def test_no_leakage_reusing_the_render_checks(draft):
    paths = draft["paths"]
    chunks = rd.load_chunk_texts(paths)
    assert cdr.draft_problems(draft["md"], chunks, [HEADLINE]) == []
    assert cdr.draft_problems(cdr.docx_text(draft["docx"]), chunks, [HEADLINE]) == []
    assert "evidence" not in draft["md"].lower().replace(cdr.POINTER.lower(), "")
    extract = json.loads((paths.stages / "02_extract.json").read_text())
    for k in ("initiatives", "challenges", "changes"):
        for item in extract[k]:
            for s in item["sentences"]:
                for c in s["citations"]:
                    if c["evidence"] not in s["text"]:
                        assert c["evidence"] not in draft["md"]


def test_csm_draft_with_a_15_word_source_run_is_refused(draft):
    paths = draft["paths"]
    out = rd.output_dir(paths)
    plan = json.loads((out / rd.PLAN_JSON).read_text())
    plan["key_business_initiatives"][0]["sentences"][0]["text"] = " ".join(CHUNK_TEXT.split()[15:30]) + "."
    (out / rd.PLAN_JSON).write_text(json.dumps(plan))
    (out / cdr.DRAFT_MD).unlink()
    (out / cdr.DRAFT_DOCX).unlink()
    with pytest.raises(rd.RenderError, match="run of 15\\+ words copied from a source chunk"):
        cdr.write_csm_draft(paths, [HEADLINE])
    assert not (out / cdr.DRAFT_MD).exists() and not (out / cdr.DRAFT_DOCX).exists()


@pytest.mark.parametrize("bad,match", [
    ("Teams pay $8 per user.", "currency amount"),
    ("Available on the Premium plan.", "plan tier"),
    ("Mr. Smith leads the work.", "personal title"),
    ("Their existing Jira instance is large.", "implies the company is a customer"),
])
def test_csm_draft_checks_catch_prices_tiers_names_and_customer_framing(bad, match):
    assert any(match in p for p in cdr.draft_problems(f"Line. {bad}", [], []))


def test_no_dates_beyond_those_in_the_plan(draft):
    plan_dates = {m for _, m in guards.date_mentions(draft["plan_md"])}
    assert {m for _, m in guards.date_mentions(draft["md"])} <= plan_dates


def test_questions_come_from_the_plan_and_appear_once(draft):
    plan = draft["plan"]
    plan_questions = {i.get("validation_question") for i in plan["key_business_initiatives"]}
    plan_questions |= {b.get("validation_question") for b in plan["key_benefits"]}
    plan_questions |= {c.get("validation_question") for c in plan["key_challenges"]["public_signals"]}
    for c in plan["success_criteria"]:
        plan_questions.add(c.get("baseline_question"))  # criterion validation questions are in plan.md only
    lines = question_lines(draft["md"])
    texts = [l.split(": ", 1)[1] for l in lines]
    assert texts and set(texts) <= plan_questions
    assert len(texts) == len(set(texts))  # each question once


def test_questions_sit_under_their_objective_criterion_or_challenge(draft):
    plan, md = draft["plan"], draft["md"]
    objectives = section_text(md, 6)
    o2 = objectives.split("- O-02 ")[1].split("\n- O-03 ")[0]
    assert "  - Validate: What is the status of I-02?" in o2 and "  - Validate: Does B-02 fit?" in o2
    c = plan["objectives"][0]["success_criteria"][0]  # the first criterion shows the (shared) fixture baseline question
    block = objectives.split(f"  - Criterion {c['id']}: ")[1]
    assert block.splitlines()[1] == f"    - Baseline: {c['baseline_question']}"
    challenges = section_text(md, 5)
    assert "- change S-01 (sourced; news headline only)\n  - Validate: What is the status of S-01?" in challenges
    criteria = section_text(md, 4)
    assert "Validate:" not in criteria and "Baseline:" not in criteria  # measures only


def test_only_baseline_questions_under_criteria(draft):
    plan, md = draft["plan"], draft["md"]
    lines = section_text(md, 6).splitlines()
    under_criteria = [l for l in lines if l.startswith("    - ")]
    assert under_criteria and all(l.startswith("    - Baseline: ") for l in under_criteria)
    baselines = {c.get("baseline_question") for c in plan["success_criteria"]}
    assert {l.split("Baseline: ", 1)[1] for l in under_criteria} <= baselines
    for c in plan["success_criteria"]:
        if c.get("validation_question") and c["validation_question"] not in {
                i.get("validation_question") for i in plan["key_business_initiatives"]} | {
                b.get("validation_question") for b in plan["key_benefits"]}:
            assert c["validation_question"] not in md  # removed from the CSM draft
            assert c["validation_question"] in draft["plan_md"]  # still in plan.md


def test_success_criteria_list_measures_with_one_target_blank(draft):
    criteria = section_text(draft["md"], 4)
    assert criteria.count(cdr.TARGET_BLANK_ONCE) == 1 and draft["md"].count(cdr.TARGET_BLANK_ONCE) == 1
    assert "- O-02-C1 Increase: number of teams using one shared view (product usage, leading; inferred)" in criteria
    assert cdr._measure({"text": "[CSM to complete: x]"}) == "[CSM to complete: x]"


def test_no_fit_questions_stay_in_the_detailed_plan(draft):
    plan, md = draft["plan"], draft["md"]
    no_fit = {b["initiative_id"]: b for b in plan["key_benefits"] if b["fit"] == "no_fit"}
    assert no_fit
    for iid, b in no_fit.items():
        assert f"What is the status of {iid}?" not in md and b["validation_question"] not in md
        assert f"What is the status of {iid}?" in draft["plan_md"]


def test_csm_blanks_kept_once_and_csm_selected_stays_null(draft):
    md = draft["md"]
    assert "Account history: [CSM to complete: account history]" in md
    assert "[CSM to complete: renewal date, ARR, status, next steps]" in md
    assert md.count(PLAN_CONFIRMATION_BLANK) == 1
    assert cdr.CUSTOMER_RAISED in md and OBJECTIVES_BLANK in md
    assert "All initiatives: Proposed, awaiting CSM review." in md
    assert {i["csm_selected"] for i in draft["plan"]["key_business_initiatives"]} == {None}
    assert not re.search(r"^- .*\[CSM to complete\]$", section_text(md, 1), re.M)  # no per-item blanks


def test_key_benefits_use_display_names_fit_labels_and_a_no_fit_count(draft):
    benefits = section_text(draft["md"], 2)
    assert "Unified technology platform (I-01): automation rules" in benefits
    assert "Cross-team delivery planning (I-02): cross-team plans" in benefits and "full fit" in benefits
    assert "No fit for 1 initiative (listed in the detailed plan)." in benefits
    assert "Restructuring" not in benefits


def test_objectives_nest_criteria_questions_and_tasks(draft):
    lines = [l for l in section_text(draft["md"], 6).splitlines()[3:] if l.strip()]
    assert lines[0] == cdr.BASELINES_ONCE
    assert lines[1].startswith("- O-01 Explore support for: Unified technology platform")
    assert any(l.startswith("  - Criterion O-01-C1: ") for l in lines)
    assert any(l.startswith("  - Task: ") and "(CSM, Day 0-30, proposed)" in l for l in lines)
    assert any(l.startswith("    - Baseline: ") for l in lines)


# ---------------------------------------------------------------- challenge grouping (config/csm_draft_rules.yaml)

def test_rules_are_loaded_from_the_config_file(tmp_path):
    rules = cdr.load_rules()
    assert rules == {"categories": {"organizational", "security_technology"},
                     "change_types": {"leadership_change", "reorganization", "restructuring_layoffs", "acquisition"}}
    f = tmp_path / "rules.yaml"
    f.write_text("challenge_categories_with_question: [competition]\nchange_types_with_question: []\n")
    assert cdr.load_rules(f) == {"categories": {"competition"}, "change_types": set()}


@pytest.mark.parametrize("fields,kept", [
    ({"category": "organizational"}, True), ({"category": "security_technology"}, True),
    ({"change_type": "acquisition"}, True), ({"change_type": "leadership_change"}, True),
    ({"category": "competition"}, False), ({"category": "other"}, False), ({"change_type": "financing"}, False),
    ({"category": "made_up_category"}, False), ({"change_type": "made_up_type"}, False), ({}, False),
    ({"category": None, "change_type": None}, False),
])
def test_unknown_or_missing_category_falls_to_context(fields, kept):
    assert cdr.keeps_question(fields, cdr.load_rules()) is kept


def challenge_section(md):
    sec = section_text(md, 5)
    return sec, sec.split(f"### {cdr.ENGAGEMENT_GROUP}")[1].split(cdr.CUSTOMER_RAISED)[0]


def test_only_kept_challenges_with_a_count_of_the_others(draft):
    md, plan = draft["md"], draft["plan"]
    sec, engaged = challenge_section(md)
    assert "Company context" not in md
    rules, fields = cdr.load_rules(), cdr.extract_fields(draft["paths"])
    signals = plan["key_challenges"]["public_signals"]
    kept = [c for c in signals if cdr.keeps_question(fields.get(c["id"], {}), rules)]
    others = len(signals) - len(kept)
    bullets = [l[2:].split(" (")[0] for l in sec.splitlines() if l.startswith("- ")]
    assert bullets == [c["title"] for c in kept]  # each kept item once, no others
    assert cdr.others_line(others) in sec
    assert re.search(rf"^{others} other public risks? (and changes are|or change is) listed in the detailed plan \(plan\.md\)\.$", sec, re.M)
    assert sec.index(cdr.others_line(others)) < sec.index(cdr.CUSTOMER_RAISED)


def test_kept_challenges_show_their_question_and_others_stay_in_plan_md(draft):
    _, engaged = challenge_section(draft["md"])
    assert "- change S-01 (sourced; news headline only)\n  - Validate: What is the status of S-01?" in engaged
    assert "challenge C-01" not in draft["md"] and "What is the status of C-01?" not in draft["md"]
    assert "What is the status of C-01?" in draft["plan_md"]


def test_others_line_wording():
    assert cdr.others_line(11) == "11 other public risks and changes are listed in the detailed plan (plan.md)."
    assert cdr.others_line(1) == "1 other public risk or change is listed in the detailed plan (plan.md)."


def test_missing_extract_fields_leave_no_kept_challenges(draft):
    paths = draft["paths"]
    (paths.stages / "02_extract.json").unlink()
    md = cdr.to_markdown(cdr.build_blocks(draft["plan"], {}, cdr.extract_fields(paths), cdr.load_rules()))
    sec, engaged = challenge_section(md)
    n = len(draft["plan"]["key_challenges"]["public_signals"])
    assert "None." in engaged and cdr.others_line(n) in sec


def test_challenge_action_line_text():
    assert cdr.ACTIONS[5] == "Validate which of these are relevant to this account and add any the customer raises."


def test_omitted_source_initiative_bullet(draft):
    paths = draft["paths"]
    out = rd.output_dir(paths)
    plan = json.loads((out / rd.PLAN_JSON).read_text())
    item = plan["key_business_initiatives"][0]
    item["text"] = rd.OMITTED_ITEM_TEXT
    item.pop("sentences", None)
    md = cdr.to_markdown(cdr.build_blocks(plan, {}, {}, cdr.load_rules()))
    assert f"- {item['title']}: [Description omitted in the CSM draft; see the detailed plan]\n" in md
    assert "source passages omitted" not in md


# ---------------------------------------------------------------- initiatives split by fit; one baselines line

def test_fit_initiatives_as_bullets_and_one_no_fit_line(draft):
    plan, md = draft["plan"], draft["md"]
    sec = section_text(md, 1)
    fit_ids = {b["initiative_id"] for b in plan["key_benefits"] if b["fit"] != "no_fit"}
    fit = [i for i in plan["key_business_initiatives"] if i["id"] in fit_ids]
    no_fit = [i for i in plan["key_business_initiatives"] if i["id"] not in fit_ids]
    bullets = [l for l in sec.splitlines() if l.startswith("- ")]
    assert [b[2:].split(":")[0] for b in bullets] == [i["title"] for i in fit]
    for i, b in zip(fit, bullets):  # fit bullets unchanged: title, first shown sentence, labels
        assert b.startswith(f"- {i['title']}: {i['sentences'][0]['text']}")
    line = next(l for l in sec.splitlines() if l.startswith(cdr.NO_FIT_PREFIX))
    assert line[len(cdr.NO_FIT_PREFIX):].split("; ") == [i["title"] for i in no_fit]
    assert sec.index(bullets[-1]) < sec.index(line)


def test_initiatives_action_line_mentions_bullets_and_no_fit_line():
    assert cdr.ACTIONS[1].startswith("Select the initiatives relevant to this account, from the bullets and from the no-fit line.")


def test_no_fit_line_absent_when_every_initiative_fits(draft):
    plan = json.loads(json.dumps(draft["plan"]))
    for b in plan["key_benefits"]:
        b["fit"] = "partial"
    assert cdr.NO_FIT_PREFIX not in cdr.to_markdown(cdr.build_blocks(plan, {}, {}, cdr.load_rules()))


def test_baselines_task_once_at_the_top_of_objectives(draft):
    md, plan = draft["md"], draft["plan"]
    sec = section_text(md, 6)
    baselines = cdr.baseline_task_text()
    assert baselines == "Agree baselines for this objective's success criteria."
    assert md.count(cdr.BASELINES_ONCE) == 1 and baselines not in md  # no per-objective copies in the CSM draft
    first_objective = sec.index("- O-01 ")
    assert sec.index(cdr.BASELINES_ONCE) < first_objective
    assert all(any(t["text"] == baselines for t in ob["tasks"]) for ob in plan["objectives"])  # plan.json keeps them
    assert draft["plan_md"].count(baselines) == len(plan["objectives"])  # and so does plan.md
